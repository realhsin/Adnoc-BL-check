import tempfile
from pathlib import Path

from docx import Document
from fpdf import FPDF

from comparison_engine.pipelines.document_pipeline import DocumentPipeline
from comparison_engine.services.report_generator import ReportGenerator


def _create_docx(path: Path, text: str) -> None:
    document = Document()
    for line in text.strip().splitlines():
        document.add_paragraph(line.strip())
    document.save(path)


def _create_pdf(path: Path, text: str) -> None:
    pdf = FPDF(unit="mm", format="A4")
    pdf.add_page()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.set_font("Helvetica", size=10)
    for line in text.strip().splitlines():
        pdf.cell(0, 6, txt=line.strip(), ln=1)
    pdf.output(str(path))


def test_extract_shipment_from_docx_text():
    raw_text = """
BL Number: BL123
Container Number: CONT001
Container Type: Dry
Origin: Dubai
Destination: Abu Dhabi
Commodity: Electronics
Quantity: 100
Gross Weight: 2000
Net Weight: 1800
Vessel: VesselA
Voyage: V001
"""
    shipment = DocumentPipeline.extract_shipment(raw_text)

    assert shipment.bl_number == "BL123"
    assert shipment.container_number == "CONT001"
    assert shipment.container_type == "Dry"
    assert shipment.origin == "Dubai"
    assert shipment.destination == "Abu Dhabi"
    assert shipment.quantity == "100"
    assert shipment.gross_weight == "2000"
    assert shipment.net_weight == "1800"


def test_compare_document_pair_with_docx_files(tmp_path: Path):
    raw_left = """
BL Number: BL123
Container Number: CONT001
Container Type: Dry
Origin: Dubai
Destination: Abu Dhabi
Commodity: Electronics
Quantity: 100
Gross Weight: 2000
Net Weight: 1800
Vessel: VesselA
Voyage: V001
"""
    raw_right = """
BL Number: BL123
Container Number: CONT001
Container Type: dry
Origin: Dubai
Destination: Abu Dhabi
Commodity: Electronics
Quantity: 100.0
Gross Weight: 2,000
Net Weight: 1800
Vessel: VesselA
Voyage: V001
"""
    left_path = tmp_path / "left.docx"
    right_path = tmp_path / "right.docx"
    _create_docx(left_path, raw_left)
    _create_docx(right_path, raw_right)

    report = DocumentPipeline.compare_document_pair(str(left_path), str(right_path))["report"]
    assert report.summary["matched_fields"] == 11
    assert report.summary["mismatched_fields"] == 0
    assert report.summary["missing_fields"] == 0
    assert report.summary["invalid_fields"] == 0


def test_compare_document_pair_with_pdf_files(tmp_path: Path):
    raw_left = """
BL Number: BL123
Container Number: CONT001
Container Type: Dry
Origin: Dubai
Destination: Abu Dhabi
Commodity: Electronics
Quantity: 100
Gross Weight: 2000
Net Weight: 1800
Vessel: VesselA
Voyage: V001
"""
    raw_right = """
BL Number: BL123
Container Number: CONT001
Container Type: dry
Origin: Dubai
Destination: Abu Dhabi
Commodity: Electronics
Quantity: 100.0
Gross Weight: 2,000
Net Weight: 1800
Vessel: VesselA
Voyage: V001
"""
    left_path = tmp_path / "left.pdf"
    right_path = tmp_path / "right.pdf"
    _create_pdf(left_path, raw_left)
    _create_pdf(right_path, raw_right)

    report = DocumentPipeline.compare_document_pair(str(left_path), str(right_path))["report"]
    assert report.summary["matched_fields"] == 11
    assert report.summary["mismatched_fields"] == 0
    assert report.summary["missing_fields"] == 0
    assert report.summary["invalid_fields"] == 0


def test_report_generator_builds_report_for_document_pipeline(tmp_path: Path):
    raw_text = """
BL Number: BL123
Container Number: CONT001
Container Type: Dry
Origin: Dubai
Destination: Abu Dhabi
Commodity: Electronics
Quantity: 100
Gross Weight: 2000
Net Weight: 1800
Vessel: VesselA
Voyage: V001
"""
    left_path = tmp_path / "left.docx"
    right_path = tmp_path / "right.docx"
    _create_docx(left_path, raw_text)
    _create_docx(right_path, raw_text)

    comparison = DocumentPipeline.compare_document_pair(str(left_path), str(right_path))
    report = ReportGenerator.build_report(comparison["report"])

    assert report["summary"]["total_fields"] == 11
    assert report["summary"]["matched_fields"] == 11
